
from __future__ import annotations

import csv
import html
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx",
}


@dataclass
class ExtractedDocument:
    text: str
    metadata: dict[str, Any]
    warnings: list[str]


class DocumentTextExtractor:
    """Domain-neutral file text extraction for runtime knowledge ingestion."""

    def supported_extensions(self) -> list[str]:
        return sorted(SUPPORTED_EXTENSIONS)

    def extract(self, path: Path, *, original_name: str | None = None, content_type: str | None = None) -> ExtractedDocument:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        warnings: list[str] = []
        text = ""
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"unsupported_file_type:{suffix or 'unknown'}")
        try:
            if suffix in {".txt", ".md", ".markdown"}:
                text = file_path.read_text(encoding="utf-8", errors="replace")
            elif suffix == ".csv":
                text = self._extract_csv(file_path)
            elif suffix == ".json":
                text = self._extract_json(file_path)
            elif suffix in {".html", ".htm"}:
                text = self._extract_html(file_path)
            elif suffix == ".pdf":
                text = self._extract_pdf(file_path, warnings)
            elif suffix == ".docx":
                text = self._extract_docx(file_path, warnings)
            elif suffix == ".xlsx":
                text = self._extract_xlsx(file_path, warnings)
            elif suffix == ".pptx":
                text = self._extract_pptx(file_path, warnings)
        except Exception as exc:
            raise ValueError(f"extract_failed:{exc.__class__.__name__}:{exc}") from exc
        normalized = self._normalize_text(text)
        if not normalized:
            warnings.append("no_extractable_text")
        return ExtractedDocument(
            text=normalized,
            metadata={
                "filename": original_name or file_path.name,
                "extension": suffix,
                "content_type": content_type or "",
                "size_bytes": file_path.stat().st_size if file_path.exists() else 0,
            },
            warnings=warnings,
        )

    def _extract_csv(self, path: Path) -> str:
        rows: list[str] = []
        raw = path.read_text(encoding="utf-8", errors="replace")
        for row in csv.reader(raw.splitlines()):
            rows.append(" | ".join(str(cell) for cell in row))
        return "\n".join(rows)

    def _extract_json(self, path: Path) -> str:
        raw = path.read_text(encoding="utf-8", errors="replace")
        try:
            obj = json.loads(raw)
        except Exception:
            return raw
        return json.dumps(obj, ensure_ascii=False, indent=2)

    def _extract_html(self, path: Path) -> str:
        raw = path.read_text(encoding="utf-8", errors="replace")
        raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        return html.unescape(raw)

    def _extract_pdf(self, path: Path, warnings: list[str]) -> str:
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception as exc:
            warnings.append(f"pdf_dependency_unavailable:{exc.__class__.__name__}")
            return ""
        reader = PdfReader(str(path))
        parts: list[str] = []
        for index, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                warnings.append(f"pdf_page_extract_failed:{index}:{exc.__class__.__name__}")
                text = ""
            if text.strip():
                parts.append(f"[page {index + 1}]\n{text}")
        return "\n\n".join(parts)

    def _extract_docx(self, path: Path, warnings: list[str]) -> str:
        try:
            from docx import Document  # type: ignore
        except Exception as exc:
            warnings.append(f"docx_dependency_unavailable:{exc.__class__.__name__}")
            return ""
        doc = Document(str(path))
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    def _extract_xlsx(self, path: Path, warnings: list[str]) -> str:
        try:
            import openpyxl  # type: ignore
        except Exception as exc:
            warnings.append(f"xlsx_dependency_unavailable:{exc.__class__.__name__}")
            return ""
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"[sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                values = ["" if cell is None else str(cell) for cell in row]
                if any(v.strip() for v in values):
                    parts.append(" | ".join(values))
        return "\n".join(parts)

    def _extract_pptx(self, path: Path, warnings: list[str]) -> str:
        parts: list[str] = []
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        try:
            with zipfile.ZipFile(path) as zf:
                slide_names = sorted(name for name in zf.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
                for index, name in enumerate(slide_names, start=1):
                    try:
                        root = ET.fromstring(zf.read(name))
                    except Exception as exc:
                        warnings.append(f"pptx_slide_extract_failed:{index}:{exc.__class__.__name__}")
                        continue
                    texts = [node.text or "" for node in root.findall(".//a:t", ns)]
                    text = " ".join(t.strip() for t in texts if t and t.strip())
                    if text:
                        parts.append(f"[slide {index}]\n{text}")
        except Exception as exc:
            warnings.append(f"pptx_extract_failed:{exc.__class__.__name__}")
        return "\n\n".join(parts)

    def _normalize_text(self, text: str) -> str:
        value = str(text or "").replace("\x00", " ")
        value = re.sub(r"[ \t\r\f\v]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()


class DocumentChunker:
    """Generic chunking independent of document domain."""

    def chunk(
        self,
        text: str,
        *,
        document_id: str | None = None,
        max_chars: int = 1200,
        overlap_chars: int = 160,
    ) -> list[dict[str, Any]]:
        cleaned = str(text or "").strip()
        if not cleaned:
            return []
        safe_max = max(300, int(max_chars or 1200))
        safe_overlap = max(0, min(int(overlap_chars or 0), safe_max // 2))
        units = self._paragraph_units(cleaned)
        chunks: list[dict[str, Any]] = []
        current = ""
        cursor = 0

        for unit in units:
            candidate = (current + "\n\n" + unit).strip() if current else unit
            if current and len(candidate) > safe_max:
                text_value = current.strip()
                start = max(0, cursor - len(text_value))
                chunks.append(self._make_chunk(text_value, len(chunks), document_id=document_id, char_start=start))
                prefix = text_value[-safe_overlap:] if safe_overlap > 0 else ""
                current = (prefix + "\n" + unit).strip() if prefix else unit
            elif len(unit) > safe_max:
                if current.strip():
                    text_value = current.strip()
                    start = max(0, cursor - len(text_value))
                    chunks.append(self._make_chunk(text_value, len(chunks), document_id=document_id, char_start=start))
                    current = ""
                for part in self._split_large_unit(unit, max_chars=safe_max, overlap_chars=safe_overlap):
                    chunks.append(self._make_chunk(part, len(chunks), document_id=document_id, char_start=cursor))
                    cursor += max(1, len(part) - safe_overlap)
                continue
            else:
                current = candidate
            cursor += len(unit) + 2

        if current.strip():
            text_value = current.strip()
            start = max(0, cursor - len(text_value))
            chunks.append(self._make_chunk(text_value, len(chunks), document_id=document_id, char_start=start))
        self._link_chunks(chunks)
        return chunks

    def _make_chunk(self, text: str, index: int, *, document_id: str | None, char_start: int) -> dict[str, Any]:
        chunk_id = f"{document_id or 'document'}_chunk_{index}"
        value = str(text or "").strip()
        return {
            "chunk_id": chunk_id,
            "chunk_index": index,
            "text": value,
            "char_count": len(value),
            "char_start": max(0, int(char_start or 0)),
            "char_end": max(0, int(char_start or 0)) + len(value),
            "token_estimate": self._token_estimate(value),
        }

    def _link_chunks(self, chunks: list[dict[str, Any]]) -> None:
        for index, chunk in enumerate(chunks):
            chunk["previous_chunk_id"] = chunks[index - 1]["chunk_id"] if index > 0 else None
            chunk["next_chunk_id"] = chunks[index + 1]["chunk_id"] if index + 1 < len(chunks) else None

    def _split_large_unit(self, text: str, *, max_chars: int, overlap_chars: int) -> list[str]:
        value = str(text or "").strip()
        if not value:
            return []
        parts: list[str] = []
        step = max(1, max_chars - overlap_chars)
        for start in range(0, len(value), step):
            part = value[start:start + max_chars].strip()
            if part:
                parts.append(part)
            if start + max_chars >= len(value):
                break
        return parts

    def _token_estimate(self, text: str) -> int:
        value = str(text or "")
        word_like = re.findall(r"[\w\-]+", value, flags=re.UNICODE)
        if word_like:
            return max(1, int(len(word_like) * 1.25))
        return max(1, len(value) // 4)

    def _paragraph_units(self, text: str) -> list[str]:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        if len(paragraphs) > 1:
            return paragraphs
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if len(lines) > 1:
            return lines
        return [text]
